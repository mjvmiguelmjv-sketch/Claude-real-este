"""Hoja de rodaje por formatos, con el formato exacto de la hoja v3 de Miguel (13 oct).

    python hoja_rodaje.py <hoja_plantilla.docx> <guiones_mes.docx> <salida.docx> "<mes>" "<nota de cabecera>"

Saca de los guiones del mes cada reel —pilar, formato, gancho escrito y lo que se dice
en cada fase— y los agrupa por montaje de cámara para grabarlos del tirón. No va en
orden de publicación. Detrás, dos partes de material: lo que hay que tener antes de
grabar cada reel (SOP-09, pasos 6 y 7) y lo que piden de fondo las historias, que no
se graban hablando pero casi todas llevan un vídeo mudo, una foto o una captura.
"""
import copy, re, sys
import docx
sys.path.insert(0, __file__.rsplit('/', 1)[0])
from docxedit import body_els, piece_code, ptext, style, W_P, qn
from piezas import table, para, cell_paras, set_p, fill_cell

W_TR, W_TC = qn('w:tr'), qn('w:tc')

BLOQUES = [
    ('Sentado, plano medio', 'Pantalla dividida y reacción comparten cámara: tú, sentado y en plano medio; la pantalla o el clip de arriba se montan después.'),
    ('De pie, en la pizarra', 'Plano medio, la pizarra a tu derecha. Se escribe en directo lo que pide cada guion.'),
    ('Espontáneo, con el móvil en mano', 'Cámara frontal y el móvil en mano, como dice cada guion: en casa, de pie y con luz de ventana, salvo donde pone otro sitio.'),
]

# Lo que no se pide como material: remite a un fondo anterior, o es diseño
NO_SE_PIDE = ('El mismo', 'La misma', 'Los mismos', 'Las mismas', 'Fondo liso', 'Igual')
# Filas con PENDIENTE que no son material de rodaje (van por otro lado)
FUERA_DE_MATERIAL = ('Destino del CTA', 'Vídeos de referencia', 'Guion del banco', 'Elección', 'Descartados',
                     'Formato', 'Visual', 'Gancho B, para probar', 'De dónde sale cada uno', 'Fondo, NOTA DE EDICIÓN',
                     'Nota de edición')


def clasifica(formato):
    # manda la primera frase: el resto describe el plano y puede nombrar otros formatos
    cabeza = formato.split('.')[0].lower()
    f = formato.lower()
    if cabeza.startswith('a cámara con objeto'):
        return (0, 'A cámara con objeto, sentado') if 'sentado' in f else (2, 'A cámara con objeto, de pie en casa')
    for clave, res in (('horizontal', (1, 'Horizontal + texto (corte de YouTube)')), ('reacción', (0, 'Reacción')),
                       ('pantalla dividida', (0, 'Pantalla dividida')), ('pizarra', (1, 'Pizarra'))):
        if cabeza.startswith(clave): return res
    if cabeza.startswith('a cámara'): return 0, 'A cámara, sentado'
    if 'reacción' in f: return 0, 'Reacción'
    if 'pantalla dividida' in f: return 0, 'Pantalla dividida'
    if 'pizarra' in f: return 1, 'Pizarra'
    if 'horizontal' in f: return 1, 'Horizontal + texto (corte de YouTube)'
    if 'espontáneo' in f or 'móvil' in f:
        donde = 'en el coche' if 'coche' in f else 'en la calle' if 'calle' in f else 'en casa'
        return 2, f'Espontáneo (móvil), {donde}'
    return 0, 'A cámara, sentado'


def reels(path):
    d = docx.Document(path); out = []; cur = None
    for el in body_els(d):
        c = piece_code(el)
        if c:
            cur = None
            if not c.startswith(('S', 'NS')):
                cur = {'codigo': c, 'titulo': ptext(el), 'ficha': {}, 'gancho': {}, 'dices': [], 'fondos': [],
                       'pendientes': [], 'sec': None}
                out.append(cur)
            continue
        if el.tag == W_P and style(el) in ('Heading1', 'Heading2'):
            cur = None
        if cur is None: continue
        if el.tag == W_P:
            t = ptext(el).strip()
            if t.startswith('Ficha de la'): cur['sec'] = 'ficha'
            elif t == 'El gancho': cur['sec'] = 'gancho'
            elif t.startswith(('FASE', 'REMATE')): cur['sec'] = 'fase'
            elif t in ('Pie de publicación', 'Triggers usados', 'Disparadores usados', 'Guion de producción'):
                cur['sec'] = None
            continue
        for tr in el.iter(W_TR):
            tcs = tr.findall(W_TC)
            if len(tcs) < 2: continue
            k = ptext(tcs[0]).strip(); v = ' '.join(cell_paras(tcs[1]))
            if cur['sec'] == 'ficha': cur['ficha'][k] = v
            elif cur['sec'] == 'gancho': cur['gancho'][k] = v
            elif cur['sec'] == 'fase' and k == 'Locución': cur['dices'].append(v)
            elif cur['sec'] == 'fase' and k.startswith('Fondo') and 'NOTA' not in k.upper():
                cur['fondos'].append((k, v))
            if 'PENDIENTE' in v:
                cur['pendientes'].append((k, v))
    return out


def historias(path):
    """Cada secuencia con el fondo de cada pantalla."""
    d = docx.Document(path); out = []; cur = None; pant = None
    for el in body_els(d):
        c = piece_code(el)
        if c:
            cur = {'codigo': c, 'titulo': ptext(el), 'fondos': []} if c.startswith(('S', 'NS')) else None
            if cur: out.append(cur)
            continue
        if el.tag == W_P and style(el) in ('Heading1', 'Heading2'):
            cur = None
        if cur is None: continue
        if el.tag == W_P:
            t = ptext(el).strip()
            if t.startswith('PANTALLA'): pant = t.split(',')[0].replace('PANTALLA', 'Pantalla')
            continue
        for tr in el.iter(W_TR):
            tcs = tr.findall(W_TC)
            if len(tcs) >= 2 and ptext(tcs[0]).strip() == 'Fondo':
                cur['fondos'].append((pant, ' '.join(cell_paras(tcs[1]))))
    return out


def indice(path):
    """{código: soporte con su día}, de la tabla del índice del Doc del mes."""
    d = docx.Document(path); out = {}
    h = next(e for e in body_els(d) if e.tag == W_P and style(e) == 'Heading2' and ptext(e).strip() == 'Índice')
    for tr in h.getnext().iter(W_TR):
        v = [ptext(tc).strip() for tc in tr.findall(W_TC)]
        if len(v) > 1 and v[0] != 'Código': out[v[0]] = v[1]
    return out


def dia(soporte):
    """«Reel · jue 12 nov» → «jue 12 nov»."""
    return soporte.split('·', 1)[-1].strip() if soporte else ''


def cabecera(tabla, textos):
    """La fila de cabecera oscura de una tabla de la plantilla, con otros textos."""
    h = copy.deepcopy(tabla.findall(W_TR)[0])
    for tc, v in zip(h.findall(W_TC), textos): fill_cell(tc, [v])
    return h


def con_cabecera(tpl, cab, filas):
    t = table(tpl, filas)
    t.insert(list(t).index(t.findall(W_TR)[0]), cabecera(tpl, cab))
    return t


def main(plantilla, guiones, salida, mes, nota):
    doc = docx.Document(plantilla)
    els = body_els(doc)
    h1, intro, resumen = els[0], els[1], els[2]
    parte, h2, h2_intro = els[3], els[4], els[5]
    h3, linea, etiqueta, ficha_t = els[6], els[7], els[8], els[9]
    gancho_t, dices_t = els[11], els[13]
    resumen_tpl = copy.deepcopy(resumen)
    parte, h2, h2_intro, h3, linea, etiqueta, ficha_t, gancho_t, dices_t = [
        copy.deepcopy(x) for x in (parte, h2, h2_intro, h3, linea, etiqueta, ficha_t, gancho_t, dices_t)]
    body = doc.element.body
    for el in els[3:]:
        if el.tag != qn('w:sectPr'): body.remove(el)

    piezas = reels(guiones)
    grupos = {}
    for p in piezas:
        b, sub = clasifica(p['ficha'].get('Formato', ''))
        p['bloque'], p['sub'] = b, sub
        grupos.setdefault(b, {}).setdefault(sub, []).append(p)

    # cabecera y tabla resumen
    set_p(h1, f'Hoja de rodaje · {mes}')
    total = len(piezas)
    set_p(intro, f'Va por formato, para grabarlo todo del tirón: no va en orden de publicación. {nota} '
                 f'{total} reels en {len(grupos)} bloques. Al grabar, una carpeta por pieza con su código (R01, N01…).')
    filas = []
    for b in sorted(grupos):
        for sub, ps in grupos[b].items():
            filas.append([[f'{b + 1} · {BLOQUES[b][0]}'], [sub], [' · '.join(p['codigo'] for p in ps)]])
    nuevo = con_cabecera(resumen_tpl, ['Bloque', 'Formato', 'Piezas, en orden de grabación'], filas)
    resumen.addprevious(nuevo); body.remove(resumen)

    last = nuevo
    def add(el):
        nonlocal last
        last.addnext(el); last = el

    # las piezas, por bloque y formato
    for n, b in enumerate(sorted(grupos)):
        add(para(parte, f'PARTE {n + 1}'))
        add(para(h2, BLOQUES[b][0]))
        add(para(h2_intro, BLOQUES[b][1]))
        for sub, ps in grupos[b].items():
            for p in ps:
                pilar = p['ficha'].get('Pilar', '').split('.')[0].split(' (')[0].strip()
                add(para(h3, p['titulo']))
                add(para(linea, f'{pilar} · {p["sub"].split(",")[0]}'))
                add(para(etiqueta, 'Ficha de la pieza'))
                add(table(ficha_t, [[['Pilar'], [pilar]], [['Formato'], [p['ficha'].get('Formato', 'PENDIENTE: sin formato en el guion')]]]))
                add(para(etiqueta, 'El gancho'))
                add(table(gancho_t, [[['Escrito'], [p['gancho'].get('Escrito', 'PENDIENTE')]]]))
                add(para(etiqueta, 'Lo que dices'))
                add(table(dices_t, [[[str(i + 1)], [t]] for i, t in enumerate(p['dices'])]))

    # lo que hay que tener antes de grabar (SOP-09, pasos 6 y 7)
    idx = indice(guiones)
    filas_m = []
    for b in sorted(grupos):
        for sub, ps in grupos[b].items():
            for p in ps:
                items, vistos = [], set()
                for k, v in p['fondos']:
                    if 'A-ROLL' in k or 'a cámara' in k or v.startswith(NO_SE_PIDE) or v in vistos: continue
                    vistos.add(v); items.append(f"{k.replace('Fondo, ', '').capitalize()}: {v}")
                for k, v in p['pendientes']:
                    if k in FUERA_DE_MATERIAL or v in vistos: continue
                    vistos.add(v); items.append(f"{k}: {v}")
                if items:
                    filas_m.append([[f"{p['codigo']} · {dia(idx.get(p['codigo'], ''))}"], items,
                                    [f'Antes de grabar el bloque {b + 1}']])
    n = len(grupos)
    add(para(parte, f'PARTE {n + 1}'))
    add(para(h2, 'Antes de grabar: lo que hay que tener'))
    add(para(h2_intro, 'Lo que sale en pantalla y no es tu cara —capturas, documentos, clips y grabaciones de pantalla— '
                       'y lo que cada guion deja pendiente. Lo consigues tú, y tiene que estar antes de la sesión de su bloque: '
                       'si el día de grabar falta, esa pieza se cae. Sacado de los fondos de cada guion (añadido por Claude, 20 oct).'))
    add(con_cabecera(resumen_tpl, ['Pieza', 'Qué hay que tener', 'Para cuándo'], filas_m))

    # las historias: lo que piden de fondo
    filas_h = []
    for h in historias(guiones):
        d_ = dia(idx.get(h['codigo'], ''))
        items = [f'{pant}: {v}' for pant, v in h['fondos'] if not v.startswith(NO_SE_PIDE)]
        if items:
            filas_h.append([[f"{h['codigo']} · {d_}"], items, [f"Dos días antes del {d_.split(' (')[0]}"]])
    add(para(parte, f'PARTE {n + 2}'))
    add(para(h2, 'Las historias: lo que se graba en mudo'))
    add(para(h2_intro, 'Las historias no se graban hablando: se diseñan. Pero casi todas llevan de fondo un vídeo mudo tuyo, '
                       'una foto o una captura, y sin eso no se pueden montar. Los vídeos mudos se graban en la misma sesión, '
                       'al final, de diez a quince segundos por toma. Los fondos lisos no están: van en el diseño (añadido por Claude, 20 oct).'))
    add(con_cabecera(resumen_tpl, ['Secuencia', 'Qué hay que tener, pantalla a pantalla', 'Para cuándo'], filas_h))

    doc.save(salida)
    print('guardado', salida, total, 'reels ·', len(filas_m), 'piezas con material ·', len(filas_h), 'secuencias con fondo')


if __name__ == '__main__':
    main(*sys.argv[1:6])
